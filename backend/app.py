import os
import json
import io
from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from flask_jwt_extended import JWTManager, create_access_token, jwt_required, get_jwt_identity
from passlib.hash import pbkdf2_sha256
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from dotenv import load_dotenv

# Load environment variables from .env file if it exists
load_dotenv()

app = Flask(__name__)
CORS(app)

# Database Configuration
# Use DATABASE_URL from environment for production (Postgres), fallback to local SQLite
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'sqlite:///marathi_notations.db')
# Fix for Render/Heroku Postgres URLs (they start with postgres:// but SQLAlchemy needs postgresql://)
if app.config['SQLALCHEMY_DATABASE_URI'].startswith("postgres://"):
    app.config['SQLALCHEMY_DATABASE_URI'] = app.config['SQLALCHEMY_DATABASE_URI'].replace("postgres://", "postgresql://", 1)

app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['JWT_SECRET_KEY'] = os.environ.get('JWT_SECRET_KEY', 'super-secret-key-change-this-in-production')
import datetime
app.config['JWT_ACCESS_TOKEN_EXPIRES'] = datetime.timedelta(days=30)

db = SQLAlchemy(app)
jwt = JWTManager(app)

with app.app_context():
    db.create_all()
    print(f" * Connected to Database: {'PostgreSQL (Neon)' if 'postgres' in app.config['SQLALCHEMY_DATABASE_URI'] else 'SQLite (Local)'}")

import datetime

# Models
class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password = db.Column(db.String(120), nullable=False)
    secret_question = db.Column(db.String(200), nullable=True)
    secret_answer = db.Column(db.String(120), nullable=True)
    notations = db.relationship('Notation', backref='user', lazy=True)

class Notation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    title = db.Column(db.String(200), nullable=False)
    taal_key = db.Column(db.String(50), nullable=False)
    data = db.Column(db.Text, nullable=False) # JSON string of rows
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.datetime.utcnow)
    updated_at = db.Column(db.DateTime, default=datetime.datetime.utcnow, onupdate=datetime.datetime.utcnow)

with app.app_context():
    db.create_all()

@app.route('/')
def home():
    return {"status": "Backend is running", "message": "Marathi Music Notation API is active"}, 200

@app.route('/ping', methods=['GET'])
def ping():
    return jsonify({"status": "ok", "message": "Backend is reachable"}), 200

# Auth Routes
@app.route('/api/signup', methods=['POST'])
def signup():
    data = request.json
    if User.query.filter_by(username=data['username']).first():
        return jsonify({"msg": "Username already exists"}), 400
    
    hashed_password = pbkdf2_sha256.hash(data['password'])
    new_user = User(
        username=data['username'], 
        password=hashed_password,
        secret_question=data.get('secret_question'),
        secret_answer=data.get('secret_answer')
    )
    db.session.add(new_user)
    db.session.commit()
    return jsonify({"msg": "User created successfully"}), 201

@app.route('/api/login', methods=['POST'])
def login():
    data = request.json
    user = User.query.filter_by(username=data['username']).first()
    if user and pbkdf2_sha256.verify(data['password'], user.password):
        access_token = create_access_token(identity=str(user.id))
        return jsonify(access_token=access_token, username=user.username), 200
    return jsonify({"msg": "Bad username or password"}), 401

@app.route('/api/recover-password', methods=['POST'])
def recover_password():
    data = request.json
    user = User.query.filter_by(username=data['username']).first()
    if not user:
        return jsonify({"msg": "User not found"}), 404
    
    if data.get('secret_answer') == user.secret_answer:
        new_hashed = pbkdf2_sha256.hash(data['new_password'])
        user.password = new_hashed
        db.session.commit()
        return jsonify({"msg": "Password reset successful"}), 200
    
    return jsonify({"msg": "Incorrect secret answer", "question": user.secret_question}), 400

@app.route('/api/get-question/<username>', methods=['GET'])
def get_question(username):
    user = User.query.filter_by(username=username).first()
    if user:
        return jsonify({"question": user.secret_question}), 200
    return jsonify({"msg": "User not found"}), 404

# Data Routes
@app.route('/api/notations', methods=['POST'])
@jwt_required()
def save_notation():
    user_id = int(get_jwt_identity())
    data = request.json
    
    notation_id = data.get('id')
    if notation_id:
        # Convert to int to be safe
        try:
            n_id = int(notation_id)
            notation = Notation.query.filter_by(id=n_id, user_id=user_id).first()
            if notation:
                notation.title = data['title']
                notation.taal_key = data['taal_key']
                notation.data = json.dumps(data['rows'])
                db.session.commit()
                return jsonify({
                    "msg": "Notation updated", 
                    "id": notation.id,
                    "updated_at": notation.updated_at.strftime('%Y-%m-%d %H:%M:%S')
                }), 200
        except (ValueError, TypeError):
            pass

    new_notation = Notation(
        title=data['title'],
        taal_key=data['taal_key'],
        data=json.dumps(data['rows']),
        user_id=user_id
    )
    db.session.add(new_notation)
    db.session.commit()
    return jsonify({
        "msg": "Notation saved", 
        "id": new_notation.id,
        "created_at": new_notation.created_at.strftime('%Y-%m-%d %H:%M:%S')
    }), 201

@app.route('/api/notations', methods=['GET'])
@jwt_required()
def get_notations():
    user_id = get_jwt_identity()
    notations = Notation.query.filter_by(user_id=user_id).order_by(Notation.updated_at.desc()).all()
    return jsonify([{
        "id": n.id,
        "title": n.title,
        "taal_key": n.taal_key,
        "rows": json.loads(n.data),
        "created_at": n.created_at.strftime('%Y-%m-%d %H:%M:%S'),
        "updated_at": n.updated_at.strftime('%Y-%m-%d %H:%M:%S')
    } for n in notations]), 200

@app.route('/api/notations/<int:id>', methods=['DELETE'])
@jwt_required()
def delete_notation(id):
    user_id = get_jwt_identity()
    notation = Notation.query.filter_by(id=id, user_id=user_id).first()
    if notation:
        db.session.delete(notation)
        db.session.commit()
        return jsonify({"msg": "Notation deleted"}), 200
    return jsonify({"msg": "Notation not found"}), 404

# Font configuration
# For deployment on Linux (Render), we look for a font file in the backend directory.
# Defaulting to a local fonts folder. User should upload a .ttf file there.
LOCAL_FONT_PATH = os.path.join(os.path.dirname(__file__), 'fonts', 'Nirmala.ttf')
# Fallback Windows path for local dev
WINDOWS_FONT_PATH = r"C:\Windows\Fonts\Nirmala.ttf"

MARATHI_FONT_PATH = LOCAL_FONT_PATH if os.path.exists(LOCAL_FONT_PATH) else WINDOWS_FONT_PATH
PDF_FONT = 'Helvetica'
if os.path.exists(MARATHI_FONT_PATH):
    try:
        pdfmetrics.registerFont(TTFont('Nirmala', MARATHI_FONT_PATH))
        PDF_FONT = 'Nirmala'
    except Exception as e:
        print(f"Error registering font: {e}")

def get_taal_header():
    """Returns the structure for the 16-beat taal header as requested."""
    row1_beats = [str(i) for i in range(1, 9)]
    row1_notes = ["धा", "धिन्", "धिन्", "धा", "धा", "धिन्", "धिन्", "धा"]
    row1_markers = ["x", "", "", "", "2", "", "", ""]
    
    row2_beats = [str(i) for i in range(9, 17)]
    row2_notes = ["धा", "तिन्", "तिन्", "ता", "ता", "धिन्", "धिन्", "धा"]
    row2_markers = ["0", "", "", "", "3", "", "", ""] # Added common markers for 9 and 13
    
    return {
        "part1": [row1_beats, row1_notes, row1_markers],
        "part2": [row2_beats, row2_notes, row2_markers]
    }

def add_taal_header_docx(doc, taal_config):
    beats = taal_config.get('beats', 16)
    bols = taal_config.get('bols', [])
    markers = taal_config.get('markers', [])
    
    # Create beat numbers row
    beat_nums = [str(i+1) for i in range(beats)]
    
    def add_table_part(data_rows):
        table = doc.add_table(rows=len(data_rows), cols=beats)
        table.style = 'Table Grid'
        for r_idx, row_data in enumerate(data_rows):
            for c_idx, val in enumerate(row_data):
                if c_idx < beats:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = str(val)
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_table_part([beat_nums, bols, markers])
    doc.add_paragraph()

def create_docx(data):
    doc = Document()
    title = data.get('title', 'Marathi Music Notation')
    doc.add_heading(title, 0)

    taal_config = data.get('taal_config', {})
    beats = taal_config.get('beats', 16)

    # Add the mandatory taal header
    add_taal_header_docx(doc, taal_config)

    for row in data.get('rows', []):
        if row['type'] == 'header':
            p = doc.add_paragraph()
            run = p.add_run(row['content'])
            run.bold = True
            run.font.size = Pt(14)
        elif row['type'] == 'notation':
            table = doc.add_table(rows=1, cols=beats)
            table.style = 'Table Grid'
            cells = row.get('cells', [''] * beats)
            for i in range(beats):
                cell = table.cell(0, i)
                cell.text = cells[i]
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

def create_pdf(data):
    file_stream = io.BytesIO()
    doc = SimpleDocTemplate(file_stream, pagesize=A4)
    styles = getSampleStyleSheet()
    
    font_size = data.get('font_size', 12)
    
    marathi_style = ParagraphStyle(
        'MarathiStyle',
        parent=styles['Normal'],
        fontName=PDF_FONT,
        fontSize=font_size,
        alignment=1 
    )
    
    elements = []
    title = data.get('title', 'Marathi Music Notation')
    elements.append(Paragraph(title, styles['Title']))
    elements.append(Spacer(1, 12))

    taal_config = data.get('taal_config', {})
    beats = taal_config.get('beats', 16)
    bols = taal_config.get('bols', [])
    markers = taal_config.get('markers', [])
    beat_nums = [str(i+1) for i in range(beats)]

    # Taal Header - Keeping grid for header but removing for notation as requested
    header_table_data = [beat_nums, bols, markers]
    col_width = (A4[0] - 80) / beats
    t = Table(header_table_data, colWidths=[col_width]*beats)
    t.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
        ('FONTNAME', (0, 0), (-1, -1), PDF_FONT),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BACKGROUND', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTSIZE', (0, 0), (-1, -1), font_size),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 12))

    for row in data.get('rows', []):
        if row['type'] == 'header':
            elements.append(Paragraph(f"<b>{row['content']}</b>", marathi_style))
            elements.append(Spacer(1, 6))
        elif row['type'] == 'notation':
            cells = row.get('cells', [''] * beats)
            table_data = [cells]
            t = Table(table_data, colWidths=[col_width]*beats)
            t.setStyle(TableStyle([
                # Removed 'GRID' to remove the cell borders as requested
                ('FONTNAME', (0, 0), (-1, -1), PDF_FONT),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('FONTSIZE', (0, 0), (-1, -1), font_size),
            ]))
            elements.append(t)
            elements.append(Spacer(1, 6))

    doc.build(elements)
    file_stream.seek(0)
    return file_stream

def create_txt(data):
    output = []
    output.append(data.get('title', 'Marathi Music Notation'))
    output.append("=" * 40)
    
    taal_config = data.get('taal_config', {})
    beats = taal_config.get('beats', 16)
    bols = taal_config.get('bols', [])
    markers = taal_config.get('markers', [])
    beat_nums = [str(i+1) for i in range(beats)]

    output.append("\t".join(beat_nums))
    output.append("\t".join(bols))
    output.append("\t".join(markers))
    output.append("-" * 40)
    
    for row in data.get('rows', []):
        if row['type'] == 'header':
            output.append(f"\n[{row['content']}]")
        elif row['type'] == 'notation':
            cells = row.get('cells', [''] * beats)
            line = " | ".join(cells)
            output.append(f"| {line} |")
    
    return "\n".join(output)

@app.route('/api/export/docx', methods=['POST'])
def export_docx():
    data = request.json
    file_stream = create_docx(data)
    return send_file(file_stream, as_attachment=True, download_name='notation.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@app.route('/api/export/pdf', methods=['POST'])
def export_pdf():
    data = request.json
    file_stream = create_pdf(data)
    return send_file(file_stream, as_attachment=True, download_name='notation.pdf', mimetype='application/pdf')

@app.route('/api/export/txt', methods=['POST'])
def export_txt():
    data = request.json
    content = create_txt(data)
    file_stream = io.BytesIO(content.encode('utf-8'))
    return send_file(file_stream, as_attachment=True, download_name='notation.txt', mimetype='text/plain')

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
